"""CSV / TSV / XLSX / XLS parser -> cleaned, typed tables for SQL (not text dumps)."""

from __future__ import annotations

import csv
import io
from pathlib import Path

import pandas as pd

from docchat.ingest.context import ParseContext
from docchat.ingest.parsers.text import read_text
from docchat.ingest.safety import check_zip_container
from docchat.ingest.tables import frame_to_table, grid_to_tables
from docchat.schemas import ParsedDocument

# Above this size openpyxl (needed for merged-cell info) gets slow; use calamine instead.
OPENPYXL_MAX_BYTES = 8_000_000


def parse_tabular(path: Path, ctx: ParseContext) -> ParsedDocument:
    parsed = ParsedDocument(doc_id=ctx.doc_id, file_name=path.name, path=path, format="tabular",
                            parser="tabular")  # fmt: skip
    suffix = path.suffix.lower()
    if suffix in (".csv", ".tsv"):
        parsed.parser = "csv"
        parsed.tables.append(_read_csv(path))
    else:
        if suffix in (".xlsx", ".xlsm"):
            check_zip_container(path)
        use_openpyxl = suffix in (".xlsx", ".xlsm") and path.stat().st_size <= OPENPYXL_MAX_BYTES
        sheets = _openpyxl_grids(path) if use_openpyxl else _calamine_grids(path)
        parsed.parser = "openpyxl" if use_openpyxl else "calamine"
        for sheet, grid in sheets:
            ctx.report(f"{path.name}: sheet {sheet}")
            tables = grid_to_tables(grid, label=f"{path.name} / {sheet}", sheet=sheet)
            if len(tables) > 1:
                for k, t in enumerate(tables, 1):
                    t.label = f"{t.label} #{k}"
            parsed.tables.extend(tables)
        parsed.page_count = len(sheets)
    if not parsed.tables:
        parsed.warnings.append("no tables found")
    return parsed


def _read_csv(path: Path):
    text = read_text(path)
    sample = text[:20000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        sep = dialect.delimiter
    except csv.Error:
        sep = "\t" if path.suffix.lower() == ".tsv" else ","
    df = pd.read_csv(io.StringIO(text), sep=sep, dtype=str, keep_default_na=False,
                     na_values=[""], skipinitialspace=True)  # fmt: skip
    table = frame_to_table(df, label=path.name, first_row=2)
    table.title = f"Rows of {path.name}"
    return table


def _openpyxl_grids(path: Path) -> list[tuple[str, list[list]]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, data_only=True)
    out = []
    try:
        for ws in wb.worksheets:
            if ws.sheet_state != "visible":
                continue
            grid = [list(r) for r in ws.iter_rows(values_only=True)]
            # Vertically merged cells: copy the top-left value down so every row is complete.
            for rng in ws.merged_cells.ranges:
                value = ws.cell(rng.min_row, rng.min_col).value
                if rng.min_col == rng.max_col:  # vertical merge (e.g. "Region" over 2 header rows)
                    for r in range(rng.min_row + 1, rng.max_row + 1):
                        if r - 1 < len(grid):
                            grid[r - 1][rng.min_col - 1] = value
            out.append((ws.title, grid))
    finally:
        wb.close()
    return out


def _calamine_grids(path: Path) -> list[tuple[str, list[list]]]:
    from python_calamine import CalamineWorkbook

    wb = CalamineWorkbook.from_path(str(path))
    return [(name, wb.get_sheet_by_name(name).to_python(skip_empty_area=False))
            for name in wb.sheet_names]  # fmt: skip
