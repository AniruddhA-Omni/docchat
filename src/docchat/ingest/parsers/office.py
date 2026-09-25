"""DOCX and PPTX parsers (python-docx / python-pptx), preserving reading order and structure."""

from __future__ import annotations

import io
from pathlib import Path

from docx import Document as open_docx
from docx.table import Table as DocxTable
from docx.text.paragraph import Paragraph as DocxParagraph
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

from docchat.ingest.context import ParseContext
from docchat.ingest.safety import check_zip_container
from docchat.ingest.tables import rows_markdown, rows_to_table
from docchat.schemas import Element, Kind, ParsedDocument

EMU_PER_POINT = 12700


# ================================================================================================
# DOCX
# ================================================================================================
def _page_breaks(p: DocxParagraph) -> int:
    xml = p._p
    return len(xml.xpath('.//w:br[@w:type="page"]')) + len(xml.xpath(".//w:lastRenderedPageBreak"))


def _heading_level(p: DocxParagraph) -> int | None:
    name = (p.style.name if p.style is not None else "") or ""
    if name == "Title":
        return 1
    if name.startswith("Heading"):
        digits = "".join(ch for ch in name if ch.isdigit())
        return int(digits) if digits else 1
    return None


def _docx_table_rows(table: DocxTable) -> list[list[str]]:
    rows = []
    for row in table.rows:
        cells, seen = [], set()
        for cell in row.cells:
            if id(cell._tc) in seen:  # horizontally merged cells repeat the same element
                continue
            seen.add(id(cell._tc))
            cells.append(cell.text.strip())
        if any(cells):
            rows.append(cells)
    return rows


def parse_docx(path: Path, ctx: ParseContext) -> ParsedDocument:
    check_zip_container(path)
    doc = open_docx(str(path))
    parsed = ParsedDocument(doc_id=ctx.doc_id, file_name=path.name, path=path, format="docx",
                            parser="python-docx")  # fmt: skip
    page, n_tables = 1, 0
    for block in doc.iter_inner_content():
        if isinstance(block, DocxParagraph):
            if block.paragraph_format.page_break_before:
                page += 1
            text = block.text.strip()
            if text:
                level = _heading_level(block)
                style = (block.style.name if block.style is not None else "") or ""
                kind = Kind.HEADING if level else Kind.LIST if "List" in style else Kind.PARAGRAPH
                parsed.elements.append(Element(kind=kind, text=text, page=page,
                                               heading_level=level, meta={"page_approx": True}))  # fmt: skip
            page += _page_breaks(block)
        elif isinstance(block, DocxTable):
            rows = _docx_table_rows(block)
            if len(rows) < 2:
                continue
            n_tables += 1
            parsed.elements.append(Element(kind=Kind.TABLE, text=rows_markdown(rows), page=page,
                                           table_rows=rows, meta={"page_approx": True,
                                                                  "table_idx": len(parsed.tables)}))  # fmt: skip
            table = rows_to_table(rows, label=f"{path.name} table {n_tables}", page=page)
            table.title = f"Table {n_tables} in {path.name}"
            parsed.tables.append(table)

    _docx_images(doc, path, ctx, parsed)
    parsed.page_count = page
    return parsed


def _docx_images(doc, path: Path, ctx: ParseContext, parsed: ParsedDocument) -> None:
    """Embedded pictures become figure elements (OCR'd); their exact page is unknown."""
    for k, shape in enumerate(doc.inline_shapes, 1):
        try:
            rid = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
            blob = doc.part.related_parts[rid].blob
        except (AttributeError, KeyError):
            continue
        parsed.elements.append(_image_figure(blob, f"img{k}", f"Image {k} in {path.name}", None,
                                             ctx, parsed))  # fmt: skip


def _image_figure(blob: bytes, stem: str, label: str, page: int | None, ctx: ParseContext,
                  parsed: ParsedDocument) -> Element:  # fmt: skip
    ctx.render_dir.mkdir(parents=True, exist_ok=True)
    image = Image.open(io.BytesIO(blob)).convert("RGB")
    out = ctx.render_dir / f"{stem}.png"
    image.save(out)
    ocr_text = ""
    if min(image.size) >= 64:
        try:
            engine = ctx.ocr()
            ocr_text = " ".join(ln.text for ln in engine.recognize(image))
            parsed.ocr_engine = engine.name
        except Exception as exc:
            parsed.warnings.append(f"{label}: OCR failed ({exc})")
    text = f"{label}." + (f" Text in image: {ocr_text}" if ocr_text else "")
    return Element(kind=Kind.FIGURE, text=text, page=page, image_path=str(out),
                   meta={"ocr_text": ocr_text})  # fmt: skip


# ================================================================================================
# PPTX
# ================================================================================================
def _iter_shapes(shapes):
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from _iter_shapes(shape.shapes)
        else:
            yield shape


def _shape_bbox(shape) -> tuple[float, float, float, float] | None:
    if shape.left is None or shape.top is None:
        return None
    x0, y0 = shape.left / EMU_PER_POINT, shape.top / EMU_PER_POINT
    return (
        x0,
        y0,
        x0 + (shape.width or 0) / EMU_PER_POINT,
        y0 + (shape.height or 0) / EMU_PER_POINT,
    )


def _is_title(shape) -> bool:
    return shape.is_placeholder and shape.placeholder_format.type in (
        PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE,
    )  # fmt: skip


def _chart_rows(chart) -> list[list[str]]:
    plot = chart.plots[0]
    series = list(plot.series)
    rows = [["Category", *[s.name or f"Series {i + 1}" for i, s in enumerate(series)]]]
    for i, category in enumerate(plot.categories):
        rows.append([str(category), *[_num(s.values[i]) for s in series]])
    return rows


def _num(v) -> str:
    if v is None:
        return ""
    return str(int(v)) if float(v).is_integer() else str(v)


def parse_pptx(path: Path, ctx: ParseContext) -> ParsedDocument:
    check_zip_container(path)
    prs = Presentation(str(path))
    parsed = ParsedDocument(doc_id=ctx.doc_id, file_name=path.name, path=path, format="pptx",
                            parser="python-pptx", page_count=len(prs.slides))  # fmt: skip
    for sno, slide in enumerate(prs.slides, 1):
        shapes = sorted(_iter_shapes(slide.shapes), key=lambda s: (s.top or 0, s.left or 0))
        for shape in shapes:
            bbox = _shape_bbox(shape)
            if shape.has_text_frame and shape.text_frame.text.strip():
                if _is_title(shape):
                    parsed.elements.append(Element(kind=Kind.HEADING, heading_level=1, page=sno,
                                                   text=shape.text_frame.text.strip(), bbox=bbox))  # fmt: skip
                    continue
                paras = [
                    p
                    for p in shape.text_frame.paragraphs
                    if "".join(r.text for r in p.runs).strip()
                ]
                bullets = len(paras) > 1
                text = "\n".join(("  " * p.level + "- " if bullets else "")
                                 + "".join(r.text for r in p.runs).strip() for p in paras)  # fmt: skip
                parsed.elements.append(Element(kind=Kind.LIST if bullets else Kind.PARAGRAPH,
                                               text=text, page=sno, bbox=bbox))  # fmt: skip
            elif shape.has_table:
                rows = [[c.text.strip() for c in r.cells] for r in shape.table.rows]
                rows = [r for r in rows if any(r)]
                if len(rows) >= 2:
                    parsed.elements.append(Element(kind=Kind.TABLE, text=rows_markdown(rows),
                                                   page=sno, bbox=bbox, table_rows=rows,
                                                   meta={"table_idx": len(parsed.tables)}))  # fmt: skip
                    t = rows_to_table(rows, label=f"{path.name} slide {sno}", page=sno)
                    t.title = f"Table on slide {sno} of {path.name}"
                    parsed.tables.append(t)
            elif getattr(shape, "has_chart", False) and shape.has_chart:
                chart = shape.chart
                title = chart.chart_title.text_frame.text if chart.has_title else ""
                rows = _chart_rows(chart)
                text = f"Chart{': ' + title if title else ''} (exact data from the slide)\n"
                parsed.elements.append(Element(kind=Kind.TABLE, text=text + rows_markdown(rows),
                                               page=sno, bbox=bbox, table_rows=rows,
                                               meta={"chart": True, "chart_title": title,
                                                     "table_idx": len(parsed.tables)}))  # fmt: skip
                t = rows_to_table(rows, label=f"{path.name} slide {sno} chart", page=sno)
                t.title = f"Chart data on slide {sno} of {path.name}" + (
                    f": {title}" if title else ""
                )
                parsed.tables.append(t)
            elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                el = _image_figure(shape.image.blob, f"s{sno:03d}_{shape.shape_id}",
                                   f"Picture on slide {sno}", sno, ctx, parsed)  # fmt: skip
                el.bbox = bbox
                parsed.elements.append(el)
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parsed.elements.append(Element(kind=Kind.NOTES, page=sno,
                                               text=f"Speaker notes (slide {sno}): {notes}"))  # fmt: skip
    return parsed
