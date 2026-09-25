"""Generate a reproducible, multi-format demo corpus for a fictional company plus a golden Q&A set.

    uv run python scripts/make_demo_corpus.py                  # -> demo_corpus/
    uv run python scripts/make_demo_corpus.py --large-pages 300  # also a big PDF for scale tests

Every fact lives in one place (the constants below or the seeded tables), and golden answers for
spreadsheet questions are computed from the same DataFrames that are written to disk.
"""

from __future__ import annotations

import argparse
import json
import random
import tempfile
from pathlib import Path

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
import pandas as pd  # noqa: E402
from docx import Document  # noqa: E402
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Font  # noqa: E402
from PIL import Image, ImageDraw, ImageFilter, ImageFont  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.chart.data import CategoryChartData  # noqa: E402
from pptx.enum.chart import XL_CHART_TYPE  # noqa: E402
from pptx.util import Inches, Pt  # noqa: E402
from reportlab.lib import colors  # noqa: E402
from reportlab.lib.pagesizes import A4  # noqa: E402
from reportlab.lib.styles import getSampleStyleSheet  # noqa: E402
from reportlab.lib.units import cm  # noqa: E402
from reportlab.platypus import (  # noqa: E402
    Image as RLImage,
)
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from docchat.eval.golden import ExpectedSource, GoldenItem, write_golden  # noqa: E402

SEED = 42
COMPANY = "Acme Robotics"

# --- Annual report facts (USD millions) --------------------------------------------------------
REGION_REVENUE = {  # region: (FY2024, FY2025)
    "North America": (18.6, 21.4),
    "Europe": (11.9, 14.2),
    "APAC": (7.6, 9.8),
    "LATAM": (3.1, 3.2),
}
QUARTERLY_REVENUE = {"Q1": 10.9, "Q2": 11.8, "Q3": 12.4, "Q4": 13.5}
TOTAL_FY2025, TOTAL_FY2024 = 48.6, 41.2
HEADCOUNT_REPORT, HEADCOUNT_HANDBOOK = 412, 398  # planted cross-document conflict
RND_SPEND = 7.3
GUIDANCE_FY2026 = (55, 58)

# --- Spreadsheet generation --------------------------------------------------------------------
PRODUCTS = {"Atlas Arm": 24_000, "Orion AMR": 35_000, "Vega Vision Kit": 4_500}
REGIONS = list(REGION_REVENUE)
TARGET_FACTORS = {"North America": 0.95, "Europe": 1.08, "APAC": 0.90, "LATAM": 1.15}

MONTHLY_ACTIVE_ROBOTS = [820, 845, 870, 910, 960, 1010, 1060, 1125, 1100, 1150, 1170, 1184]
MONTHS = ["Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"]
UNITS_SHIPPED = {"Atlas Arm": 1240, "Orion AMR": 860, "Vega Vision Kit": 2310}


# =============================================================================================
# helpers
# =============================================================================================
def _font(size: int, bold: bool = False, serif: bool = False) -> ImageFont.ImageFont:
    names = (
        ["timesbd.ttf", "DejaVuSerif-Bold.ttf"] if serif and bold
        else ["times.ttf", "DejaVuSerif.ttf"] if serif
        else ["arialbd.ttf", "DejaVuSans-Bold.ttf"] if bold
        else ["arial.ttf", "DejaVuSans.ttf"]
    )  # fmt: skip
    for name in names:
        for folder in ("C:/Windows/Fonts", "/usr/share/fonts/truetype/dejavu", ""):
            try:
                return ImageFont.truetype(str(Path(folder) / name) if folder else name, size)
            except OSError:
                continue
    return ImageFont.load_default(size=size)


def _scan_effect(img: Image.Image, angle: float, seed: int) -> Image.Image:
    """Make a clean render look like a skewed, noisy office scan."""
    rng = np.random.default_rng(seed)
    img = img.convert("L").rotate(angle, resample=Image.BICUBIC, expand=True, fillcolor=255)
    img = img.filter(ImageFilter.GaussianBlur(0.7))
    arr = np.asarray(img).astype(np.int16)
    arr += rng.normal(0, 10, arr.shape).astype(np.int16)
    speckles = rng.random(arr.shape) < 0.0015
    arr[speckles] = rng.integers(0, 120, speckles.sum())
    return Image.fromarray(np.clip(arr, 0, 255).astype(np.uint8))


def _money(x: float) -> str:
    return f"${x:,.0f}"


# =============================================================================================
# 1. Annual report PDF (10 pages, table, embedded chart, footnotes, page footers)
# =============================================================================================
def make_annual_report(out: Path, tmp: Path) -> None:
    styles = getSampleStyleSheet()
    h1, body = styles["Heading1"], styles["BodyText"]
    small = styles["Italic"]
    small.fontSize = 8

    chart = tmp / "quarterly.png"
    fig, ax = plt.subplots(figsize=(6, 3.2), dpi=150)
    bars = ax.bar(list(QUARTERLY_REVENUE), list(QUARTERLY_REVENUE.values()), color="#2563eb")
    ax.bar_label(bars, fmt="%.1f")
    ax.set_ylabel("Revenue (USD M)")
    ax.set_title("Figure 1 - Quarterly revenue FY2025")
    ax.set_ylim(0, 16)
    fig.tight_layout()
    fig.savefig(chart)
    plt.close(fig)

    def table(rows, widths=None):
        t = Table(rows, colWidths=widths)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e3a8a")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
        ]))  # fmt: skip
        return t

    region_rows = [["Region", "FY2024 (USD M)", "FY2025 (USD M)", "YoY growth"]]
    for region, (prev, cur) in REGION_REVENUE.items():
        region_rows.append([region, f"{prev:.1f}", f"{cur:.1f}", f"{(cur / prev - 1) * 100:.1f}%"])
    region_rows.append(["Total", f"{TOTAL_FY2024:.1f}", f"{TOTAL_FY2025:.1f}",
                        f"{(TOTAL_FY2025 / TOTAL_FY2024 - 1) * 100:.1f}%"])  # fmt: skip

    story = [
        # p1
        Paragraph(f"{COMPANY} - Annual Report FY2025", styles["Title"]),
        Paragraph("1. Executive summary", h1),
        Paragraph(
            f"FY2025 was a record year for {COMPANY}. Total revenue reached USD {TOTAL_FY2025} "
            f"million, up {(TOTAL_FY2025 / TOTAL_FY2024 - 1) * 100:.0f}% from USD {TOTAL_FY2024} "
            "million in FY2024. Growth was led by APAC and by strong adoption of the Vega Vision "
            "Kit. Gross margin improved to 58% on the back of lower component costs.<super>1</super>",
            body,
        ),
        Paragraph(
            "The board approved a dividend-free reinvestment policy for FY2026, prioritising the "
            "Orion AMR v2 launch and expansion of the service organisation.",
            body,
        ),
        PageBreak(),
        # p2
        Paragraph("2. Revenue by region", h1),
        Paragraph("Table 1 shows revenue by sales region. All figures are audited.", body),
        Spacer(1, 0.3 * cm),
        table(region_rows, [5 * cm, 3.5 * cm, 3.5 * cm, 3 * cm]),
        Spacer(1, 0.3 * cm),
        Paragraph("APAC was the fastest growing region, driven by new logistics customers.", body),
        PageBreak(),
        # p3
        Paragraph("3. Quarterly performance", h1),
        Paragraph(
            "Revenue grew sequentially in every quarter of FY2025 (Figure 1). The fourth quarter "
            "benefited from year-end warehouse automation budgets.",
            body,
        ),
        RLImage(str(chart), width=15 * cm, height=8 * cm),
        PageBreak(),
        # p4
        Paragraph("4. Product portfolio", h1),
        table(
            [
                ["Product", "Category", "List price (USD)"],
                ["Atlas Arm", "Collaborative robot arm", f"{PRODUCTS['Atlas Arm']:,}"],
                ["Orion AMR", "Autonomous mobile robot", f"{PRODUCTS['Orion AMR']:,}"],
                ["Vega Vision Kit", "Machine-vision add-on", f"{PRODUCTS['Vega Vision Kit']:,}"],
            ],
            [4.5 * cm, 6 * cm, 4 * cm],
        ),  # fmt: skip
        Spacer(1, 0.3 * cm),
        Paragraph(
            "The Vega Vision Kit is sold both standalone and bundled with Atlas Arm. Orion AMR v2 "
            "is in development; see the separate product requirements document.",
            body,
        ),
        PageBreak(),
        # p5
        Paragraph("5. People", h1),
        Paragraph(
            f"At fiscal year end {COMPANY} employed {HEADCOUNT_REPORT} people across offices in "
            "Austin, Berlin and Pune. Voluntary attrition was 9.5%, down from 12.1% in FY2024.",
            body,
        ),
        PageBreak(),
        # p6
        Paragraph("6. Research and development", h1),
        Paragraph(
            f"R&amp;D spending was USD {RND_SPEND} million, or "
            f"{RND_SPEND / TOTAL_FY2025 * 100:.0f}% of revenue. Key programmes were Orion AMR v2 "
            "navigation, on-device defect detection for Vega, and a new safety-rated controller.",
            body,
        ),
        PageBreak(),
        # p7
        Paragraph("7. Principal risks", h1),
        Paragraph(
            "Supply concentration: battery cells for Orion come from a single supplier. "
            "Currency: roughly a third of revenue is billed in EUR. "
            "Talent: competition for robotics engineers remains intense in all three hubs.",
            body,
        ),
        PageBreak(),
        # p8
        Paragraph("8. Outlook", h1),
        Paragraph(
            f"For FY2026 management guides revenue of USD {GUIDANCE_FY2026[0]}-"
            f"{GUIDANCE_FY2026[1]} million, assuming Orion AMR v2 ships in the second half.",
            body,
        ),
        PageBreak(),
        # p9
        Paragraph("9. Sustainability", h1),
        Paragraph(
            "Energy use per robot manufactured fell 12% year on year. 64% of packaging is now "
            "recycled cardboard.",
            body,
        ),
        PageBreak(),
        # p10
        Paragraph("10. Appendix and notes", h1),
        Paragraph(
            "<super>1</super> Gross margin excludes one-off restructuring costs of USD 0.4 "
            "million related to the consolidation of the Berlin warehouse.",
            small,
        ),
        Paragraph("Figures are rounded to one decimal place.", small),
    ]

    def footer(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.drawString(2 * cm, 1.2 * cm, f"{COMPANY} - Annual Report FY2025")
        canvas.drawRightString(A4[0] - 2 * cm, 1.2 * cm, f"Page {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(str(out), pagesize=A4, title=f"{COMPANY} Annual Report FY2025")
    doc.build(story, onFirstPage=footer, onLaterPages=footer)


# =============================================================================================
# 2. Scanned (image-only) PDF and scanned invoice PNG -> OCR required
# =============================================================================================
def _render_page(lines: list[tuple[str, int, bool]], size=(1654, 2339)) -> Image.Image:
    img = Image.new("L", size, 255)
    draw = ImageDraw.Draw(img)
    y = 160
    for text, font_size, bold in lines:
        if text:
            draw.text((150, y), text, fill=20, font=_font(font_size, bold=bold, serif=True))
        y += int(font_size * 1.9)
    return img


def make_scanned_agreement(out: Path) -> None:
    lines = [
        ("SUPPLY AGREEMENT", 64, True),
        ("", 30, False),
        (f"between {COMPANY} Ltd. (the Buyer)", 38, False),
        ("and Kestrel Components GmbH (the Supplier)", 38, False),
        ("", 30, False),
        ("1. Effective date: 14 March 2025", 38, False),
        ("2. Term: 24 months from the effective date", 38, False),
        ("3. Scope: supply of battery cells for Orion AMR", 38, False),
        ("4. Total contract value: USD 1,250,000", 38, True),
        ("5. Payment terms: Net 45 days", 38, False),
        ("6. Governing law: Germany", 38, False),
        ("", 30, False),
        ("Signed for the Buyer: ____________________", 38, False),
        ("Signed for the Supplier: __________________", 38, False),
    ]
    _scan_effect(_render_page(lines), angle=1.3, seed=SEED).save(out, "PDF", resolution=200.0)


def make_invoice_scan(out: Path) -> None:
    lines = [
        ("INVOICE  INV-2025-0917", 58, True),
        ("Kestrel Components GmbH", 36, False),
        (f"Bill to: {COMPANY} Ltd.", 36, False),
        ("Invoice date: 2025-09-15      Due date: 2025-10-15", 34, False),
        ("", 20, False),
        ("Battery cell pack BCP-48   x 30   @ 550.00   16,500.00", 32, False),
        ("Thermal pads TP-2          x 60   @  12.50      750.00", 32, False),
        ("Freight and insurance                          1,200.00", 32, False),
        ("", 20, False),
        ("TOTAL DUE: USD 18,450.00", 46, True),
    ]
    _scan_effect(_render_page(lines, size=(1400, 1300)), angle=-2.0, seed=SEED + 1).save(out)


# =============================================================================================
# 3. Images: chart PNG and dashboard screenshot
# =============================================================================================
def make_active_robots_chart(out: Path) -> None:
    fig, ax = plt.subplots(figsize=(8, 4), dpi=140)
    ax.plot(MONTHS, MONTHLY_ACTIVE_ROBOTS, marker="o", color="#059669")
    for m, v in zip(MONTHS, MONTHLY_ACTIVE_ROBOTS, strict=True):
        ax.annotate(str(v), (m, v), textcoords="offset points", xytext=(0, 7), ha="center",
                    fontsize=8)  # fmt: skip
    ax.set_title("Monthly active robots in the field - 2025")
    ax.set_ylabel("Active robots")
    ax.grid(alpha=0.3)
    fig.tight_layout()
    fig.savefig(out)
    plt.close(fig)


def make_dashboard(out: Path) -> None:
    img = Image.new("RGB", (1400, 820), "#0f172a")
    d = ImageDraw.Draw(img)
    d.rectangle((0, 0, 1400, 90), fill="#1e293b")
    d.text((40, 25), "Ops Dashboard - Robot Fleet", fill="white", font=_font(38, bold=True))
    tiles = [
        ("Fleet uptime", "99.2%"), ("Open incidents", "7"),
        ("Avg response time", "42 min"), ("Robots online", "1,184"),
    ]  # fmt: skip
    for i, (label, value) in enumerate(tiles):
        x = 40 + i * 340
        d.rounded_rectangle((x, 140, x + 310, 360), radius=18, fill="#1e293b")
        d.text((x + 24, 165), label, fill="#94a3b8", font=_font(28))
        d.text((x + 24, 230), value, fill="#f8fafc", font=_font(72, bold=True))
    d.text((40, 420), "Incidents by product (last 30 days)", fill="white", font=_font(30))
    for i, (product, n) in enumerate([("Orion AMR", 4), ("Atlas Arm", 2), ("Vega Vision Kit", 1)]):
        y = 480 + i * 90
        d.text((40, y + 10), product, fill="#cbd5e1", font=_font(28))
        d.rectangle((330, y, 330 + n * 150, y + 55), fill="#f97316")
        d.text((345 + n * 150, y + 10), str(n), fill="white", font=_font(28, bold=True))
    img.save(out)


# =============================================================================================
# 4. DOCX product requirements (headings, table, explicit page breaks)
# =============================================================================================
def make_prd(out: Path) -> None:
    doc = Document()
    doc.add_heading("Product Requirements - Orion AMR v2", 0)
    doc.add_heading("1. Overview", 1)
    doc.add_paragraph(
        "Orion AMR v2 is the second generation of our autonomous mobile robot for warehouse "
        "intralogistics. It targets mid-size distribution centres that need to move totes and "
        "pallets between picking zones without fixed infrastructure."
    )
    doc.add_paragraph("Target launch: Q3 2026. Target list price: USD 38,000 per unit.")
    doc.add_page_break()

    doc.add_heading("2. Functional requirements", 1)
    reqs = [
        ("R-01", "Map a 20,000 m2 facility in under 2 hours", "Must"),
        ("R-02", "Dynamic obstacle avoidance for people and forklifts", "Must"),
        ("R-03", "Integrate with WMS via REST and MQTT", "Must"),
        ("R-04", "Automatic docking and charging", "Must"),
        ("R-05", "Fleet coordination for up to 100 robots", "Should"),
        ("R-06", "Elevator integration", "Could"),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Grid Accent 1"
    for cell, text in zip(table.rows[0].cells, ["ID", "Requirement", "Priority"], strict=True):
        cell.text = text
    for row in reqs:
        cells = table.add_row().cells
        for cell, text in zip(cells, row, strict=True):
            cell.text = text
    doc.add_page_break()

    doc.add_heading("3. Performance requirements", 1)
    doc.add_heading("3.1 Payload and speed", 2)
    doc.add_paragraph("Rated payload: 150 kg. Maximum speed: 2.0 m/s in open aisles.")
    doc.add_heading("3.2 Battery", 2)
    doc.add_paragraph(
        "Runtime of at least 10 hours per charge at rated payload; full charge in 90 minutes."
    )
    doc.add_heading("4. Open questions", 1)
    doc.add_paragraph("Battery cell sourcing is single-supplier; a second source is under review.")
    doc.save(out)


# =============================================================================================
# 5. PPTX board update (native chart, table, speaker notes)
# =============================================================================================
def make_board_deck(out: Path) -> None:
    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "Q4 2025 Board Update"
    s1.placeholders[1].text = f"{COMPANY} - confidential draft for the board"

    s2 = prs.slides.add_slide(prs.slide_layouts[5])
    s2.shapes.title.text = "Key metrics FY2025"
    rows = [("Metric", "Value"), ("Net promoter score", "47"), ("Gross margin", "58%"),
            ("Customer churn", "3.1%"), ("Active customers", "326")]  # fmt: skip
    tbl = s2.shapes.add_table(len(rows), 2, Inches(1.5), Inches(1.8), Inches(7), Inches(3)).table
    for r, (a, b) in enumerate(rows):
        tbl.cell(r, 0).text, tbl.cell(r, 1).text = a, b

    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = "Units shipped by product - FY2025"
    data = CategoryChartData()
    data.categories = list(UNITS_SHIPPED)
    data.add_series("Units shipped", list(UNITS_SHIPPED.values()))
    chart = s3.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.6), Inches(8), Inches(5), data
    ).chart
    chart.has_legend = False
    chart.plots[0].has_data_labels = True

    s4 = prs.slides.add_slide(prs.slide_layouts[1])
    s4.shapes.title.text = "Priorities for 2026"
    body = s4.placeholders[1].text_frame
    body.text = "Ship Orion AMR v2 (now Q3 2026)"
    for line in ["Second battery supplier", "Expand service team in APAC",
                 "Vega on-device defect detection GA"]:  # fmt: skip
        body.add_paragraph().text = line
    for p in body.paragraphs:
        p.font.size = Pt(24)
    s4.notes_slide.notes_text_frame.text = (
        "CEO to mention that the new Singapore office opens in May 2026 and will host the APAC "
        "service hub."
    )
    prs.save(out)


# =============================================================================================
# 6. Markdown handbook and TXT meeting notes
# =============================================================================================
def make_handbook(out: Path) -> None:
    out.write_text(
        f"""# {COMPANY} Engineering Handbook

## About us
{COMPANY} has {HEADCOUNT_HANDBOOK} employees across three offices (Austin, Berlin, Pune).

## On-call
Primary on-call rotates weekly. Handover happens on Mondays at 10:00 CET.
Pages must be acknowledged within 15 minutes.

## Code review
Every change needs one approving review. Changes touching the safety controller need two.

## Deployments
Production deploys are frozen from December 20 to January 5.
Rollbacks are done with `acmectl rollback <service>`.
""",
        encoding="utf-8",
    )


def make_meeting_notes(out: Path) -> None:
    out.write_text(
        """Engineering leadership sync - 2025-11-03
Attendees: Dana (VP Eng), Luis (Robotics), Priya (Platform), Tom (Supply chain)

Decisions
1. Orion AMR v2 launch moves from Q2 2026 to Q3 2026 because the battery supplier
   (Kestrel Components) pushed first samples by ten weeks.
2. CI moves to self-hosted runners by the end of January 2026 to cut build costs.
3. Vega defect detection will ship as a paid add-on.

Action items
- Tom: qualify a second battery cell supplier by 2026-02-15.
- Priya: publish the CI migration runbook.
- Luis: re-plan the Orion v2 test campaign.
""",
        encoding="utf-8",
    )


# =============================================================================================
# 7. Code file with a subtle documented-vs-actual bug
# =============================================================================================
def make_code(out: Path) -> None:
    out.write_text(
        '''"""Inventory service used by the Acme spare-parts portal."""

from dataclasses import dataclass


class InsufficientStockError(Exception):
    """Raised when a removal would make stock negative."""


@dataclass
class Item:
    sku: str
    quantity: int
    reorder_point: int


class InventoryService:
    def __init__(self) -> None:
        self._items: dict[str, Item] = {}

    def add_stock(self, sku: str, qty: int, reorder_point: int = 10) -> Item:
        item = self._items.setdefault(sku, Item(sku, 0, reorder_point))
        item.quantity += qty
        return item

    def remove_stock(self, sku: str, qty: int) -> Item:
        item = self._items[sku]
        if qty > item.quantity:
            raise InsufficientStockError(f"{sku}: requested {qty}, have {item.quantity}")
        item.quantity -= qty
        return item

    def needs_reorder(self, sku: str) -> bool:
        item = self._items[sku]
        return item.quantity <= item.reorder_point

    def list_items(self, page: int = 1, page_size: int = 20) -> list[Item]:
        """Return one page of items. `page` is 1-based."""
        items = sorted(self._items.values(), key=lambda i: i.sku)
        start = page * page_size
        return items[start : start + page_size]
''',
        encoding="utf-8",
    )


# =============================================================================================
# 8. Spreadsheets: XLSX (clean orders, messy summary with merged headers + totals, targets), CSV
# =============================================================================================
def build_orders(rng: random.Random) -> pd.DataFrame:
    rows = []
    for n in range(1, 401):
        product = rng.choices(list(PRODUCTS), weights=[0.3, 0.2, 0.5])[0]
        units = rng.randint(1, 20) if product == "Vega Vision Kit" else rng.randint(1, 8)
        rows.append({
            "order_id": f"SO-{n:04d}",
            "order_date": pd.Timestamp("2025-01-01") + pd.Timedelta(days=rng.randint(0, 364)),
            "region": rng.choices(REGIONS, weights=[0.4, 0.3, 0.2, 0.1])[0],
            "product": product,
            "units": units,
            "unit_price": PRODUCTS[product],
        })  # fmt: skip
    df = pd.DataFrame(rows).sort_values("order_date", kind="stable").reset_index(drop=True)
    df["revenue"] = df["units"] * df["unit_price"]
    return df


def build_tickets(rng: random.Random) -> pd.DataFrame:
    rows = []
    for n in range(1, 251):
        severity = rng.choices(["low", "medium", "high", "critical"], weights=[4, 3, 2, 1])[0]
        status = rng.choices(["closed", "open"], weights=[4, 1])[0]
        base = {"low": 30, "medium": 18, "high": 9, "critical": 4}[severity]
        rows.append({
            "ticket_id": f"T-{n:04d}",
            "created_at": (pd.Timestamp("2025-01-01")
                           + pd.Timedelta(hours=rng.randint(0, 364 * 24))).isoformat(),
            "product": rng.choice(list(PRODUCTS)),
            "severity": severity,
            "status": status,
            "resolution_hours": round(rng.uniform(0.5, 2.0) * base, 1) if status == "closed"
            else None,
        })  # fmt: skip
    return pd.DataFrame(rows)


def write_sales_workbook(out: Path, orders: pd.DataFrame, targets: dict[str, int]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Orders"
    ws.append(list(orders.columns))
    for rec in orders.itertuples(index=False):
        ws.append([rec.order_id, rec.order_date.date(), rec.region, rec.product, rec.units,
                   rec.unit_price, rec.revenue])  # fmt: skip

    # A deliberately messy "report style" sheet: title rows, merged 2-level header, totals row.
    ss = wb.create_sheet("Regional Summary")
    ss["A1"] = f"{COMPANY} - Regional Summary FY2025"
    ss["A1"].font = Font(bold=True, size=14)
    ss["A2"] = "Source: order system export, unaudited"
    ss.append([])
    ss.append(["Region", "H1", None, "H2", None])
    ss.merge_cells("B4:C4")
    ss.merge_cells("D4:E4")
    ss.append([None, "Revenue (USD)", "Orders", "Revenue (USD)", "Orders"])
    ss.merge_cells("A4:A5")
    half = np.where(orders["order_date"].dt.month <= 6, "H1", "H2")
    for region in REGIONS:
        sub = orders[orders["region"] == region]
        h = half[orders["region"] == region]
        ss.append([region, int(sub[h == "H1"]["revenue"].sum()), int((h == "H1").sum()),
                   int(sub[h == "H2"]["revenue"].sum()), int((h == "H2").sum())])  # fmt: skip
    ss.append(["Total", int(orders[half == "H1"]["revenue"].sum()), int((half == "H1").sum()),
               int(orders[half == "H2"]["revenue"].sum()), int((half == "H2").sum())])  # fmt: skip
    for cell in ss[4] + ss[5]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")

    ts = wb.create_sheet("Targets")
    ts.append(["region", "target_revenue_fy2025"])
    for region, target in targets.items():
        ts.append([region, target])
    wb.save(out)


# =============================================================================================
# 9. Optional large PDF (scale / needle-in-a-haystack)
# =============================================================================================
NEEDLE_PAGE_FRACTION = 0.79


def make_large_policy_pdf(out: Path, pages: int) -> int:
    styles = getSampleStyleSheet()
    rng = random.Random(SEED)
    topics = ["travel", "procurement", "security", "expenses", "equipment", "remote work"]
    needle_page = max(1, int(pages * NEEDLE_PAGE_FRACTION))
    story = []
    for p in range(1, pages + 1):
        topic = topics[p % len(topics)]
        story.append(Paragraph(f"Policy section {p}: {topic}", styles["Heading2"]))
        for _ in range(6):
            words = " ".join(rng.choice(["employees", "must", "submit", "approval", "within",
                                         "business", "days", "manager", "records", "policy",
                                         "request", "exceptions", "finance", "review"])
                             for _ in range(60))  # fmt: skip
            story.append(Paragraph(words.capitalize() + ".", styles["BodyText"]))
        if p == needle_page:
            story.append(Paragraph(
                "Special rule: the daily meal allowance for business trips to Tokyo is JPY 9,500.",
                styles["BodyText"],
            ))  # fmt: skip
        story.append(PageBreak())
    SimpleDocTemplate(str(out), pagesize=A4).build(story)
    return needle_page


# =============================================================================================
# golden set
# =============================================================================================
def build_golden(orders: pd.DataFrame, tickets: pd.DataFrame, targets: dict[str, int],
                 needle_page: int | None) -> list[GoldenItem]:  # fmt: skip
    G, S = GoldenItem, ExpectedSource
    rep, prd, deck = "acme_annual_report_2025.pdf", "orion_amr_v2_prd.docx", "board_update_q4.pptx"
    xlsx, csv = "sales_2025.xlsx", "support_tickets.csv"

    eu_rev = int(orders.loc[orders.region == "Europe", "revenue"].sum())
    apac_units = orders[orders.region == "APAC"].groupby("product")["units"].sum()
    na_q4 = int(((orders.region == "North America") & (orders.order_date.dt.quarter == 4)).sum())
    actual = orders.groupby("region")["revenue"].sum()
    missed = sorted(r for r in REGIONS if actual[r] < targets[r])
    latam_h2 = int(orders[(orders.region == "LATAM") & (orders.order_date.dt.month > 6)]
                   ["revenue"].sum())  # fmt: skip
    orion_high_open = int(((tickets["product"] == "Orion AMR") & (tickets.severity == "high")
                           & (tickets.status == "open")).sum())  # fmt: skip
    crit_avg = round(
        float(tickets.loc[tickets.severity == "critical", "resolution_hours"].mean()), 1
    )
    n_crit = int((tickets.severity == "critical").sum())

    items = [
        # lookups
        G(id="lk-01", type="lookup", question="What was Acme Robotics' total revenue in FY2025?",
          reference="USD 48.6 million, up 18% from USD 41.2 million.",
          must_cite=[S(file=rep, page=1)], expected_numbers=["48.6"]),
        G(id="lk-02", type="lookup", question="What revenue guidance did management give for FY2026?",
          reference="USD 55-58 million.", must_cite=[S(file=rep, page=8)],
          expected_numbers=["55", "58"]),
        G(id="lk-03", type="lookup", question="How much did Acme spend on R&D in FY2025?",
          reference="USD 7.3 million, 15% of revenue.", must_cite=[S(file=rep, page=6)],
          expected_numbers=["7.3"]),
        G(id="lk-04", type="lookup", question="What payload must the Orion AMR v2 support?",
          reference="150 kg rated payload.", must_cite=[S(file=prd, page=3)],
          expected_numbers=["150"]),
        G(id="lk-05", type="lookup", question="When are production deploys frozen?",
          reference="From December 20 to January 5.", must_cite=[S(file="engineering_handbook.md")],
          expected_keywords=["December 20", "January 5"]),
        G(id="lk-06", type="lookup", question="What does footnote 1 in the annual report say?",
          reference="Gross margin excludes USD 0.4M one-off restructuring costs (Berlin warehouse).",
          must_cite=[S(file=rep, page=10)], expected_numbers=["0.4"]),
        G(id="lk-07", type="lookup", question="When does the Singapore office open?",
          reference="May 2026 (board deck speaker notes).", must_cite=[S(file=deck, page=4)],
          expected_keywords=["May 2026"]),
        G(id="lk-08", type="lookup", question="What was the net promoter score in FY2025?",
          reference="47.", must_cite=[S(file=deck, page=2)], expected_numbers=["47"]),
        # tables in documents
        G(id="tb-01", type="table", question="Which region grew fastest in FY2025 and by how much?",
          reference="APAC, +28.9% (7.6 -> 9.8 USD M).", must_cite=[S(file=rep, page=2)],
          expected_numbers=["28.9"], expected_keywords=["APAC"]),
        G(id="tb-02", type="table", question="How many Vega Vision Kit units were shipped in FY2025?",
          reference="2,310 units (board deck chart).", must_cite=[S(file=deck, page=3)],
          expected_numbers=["2310"]),
        # spreadsheets
        G(id="xl-01", type="table", question="What is the total revenue for Europe in the Orders sheet?",
          reference=_money(eu_rev), must_cite=[S(file=xlsx)], expected_numbers=[str(eu_rev)]),
        G(id="xl-02", type="table", question="Which product sold the most units in APAC?",
          reference=f"{apac_units.idxmax()} ({int(apac_units.max())} units).",
          must_cite=[S(file=xlsx)], expected_keywords=[apac_units.idxmax()]),
        G(id="xl-03", type="table",
          question="How many orders were placed in North America in Q4 2025?",
          reference=str(na_q4), must_cite=[S(file=xlsx)], expected_numbers=[str(na_q4)]),
        G(id="xl-04", type="table", question="Which regions missed their FY2025 revenue target?",
          reference=", ".join(missed), must_cite=[S(file=xlsx)], expected_keywords=missed),
        G(id="xl-05", type="table",
          question="What was LATAM's H2 revenue according to the Regional Summary sheet?",
          reference=_money(latam_h2), must_cite=[S(file=xlsx)], expected_numbers=[str(latam_h2)]),
        G(id="csv-01", type="table", question="How many high-severity Orion AMR tickets are still open?",
          reference=str(orion_high_open), must_cite=[S(file=csv)],
          expected_numbers=[str(orion_high_open)]),
        G(id="csv-02", type="table",
          question="What is the average resolution time in hours for critical tickets?",
          reference=f"{crit_avg} hours", must_cite=[S(file=csv)], expected_numbers=[str(crit_avg)]),
        # OCR
        G(id="ocr-01", type="ocr",
          question="What is the total contract value of the Kestrel supply agreement?",
          reference="USD 1,250,000.", must_cite=[S(file="supplier_agreement_scanned.pdf", page=1)],
          expected_numbers=["1250000"]),
        G(id="ocr-02", type="ocr", question="What is the total due on invoice INV-2025-0917?",
          reference="USD 18,450.00, due 2025-10-15.", must_cite=[S(file="invoice_scan.png")],
          expected_numbers=["18450"]),
        # vision
        G(id="vis-01", type="vision",
          question="According to Figure 1 in the annual report, what was Q3 revenue?",
          reference="USD 12.4 million.", must_cite=[S(file=rep, page=3)], expected_numbers=["12.4"]),
        G(id="vis-02", type="vision",
          question="In which month did the number of monthly active robots decrease?",
          reference="September (1,125 in Aug -> 1,100 in Sep).",
          must_cite=[S(file="monthly_active_robots.png")], expected_keywords=["Sep"]),
        G(id="vis-03", type="vision", question="What fleet uptime does the ops dashboard show?",
          reference="99.2%.", must_cite=[S(file="ops_dashboard.png")], expected_numbers=["99.2"]),
        # code
        G(id="code-01", type="code",
          question="Is there a bug in the pagination of inventory_service.py?",
          reference="Yes: list_items documents page as 1-based but computes start = page * "
                    "page_size, so page 1 skips the first page of items.",
          must_cite=[S(file="inventory_service.py")], expected_keywords=["list_items"]),
        G(id="code-02", type="code",
          question="What happens in remove_stock when more units are requested than are in stock?",
          reference="It raises InsufficientStockError.", must_cite=[S(file="inventory_service.py")],
          expected_keywords=["InsufficientStockError"]),
        # cross-document
        G(id="xd-01", type="cross_doc", question="How many employees does Acme Robotics have?",
          reference="Sources disagree: 412 (annual report, p.5) vs 398 (engineering handbook).",
          must_cite=[S(file=rep, page=5), S(file="engineering_handbook.md")],
          expected_numbers=["412", "398"]),
        G(id="xd-02", type="cross_doc", question="When will Orion AMR v2 launch, and why did it slip?",
          reference="Q3 2026 (PRD); slipped from Q2 because the battery supplier delayed samples "
                    "by ten weeks (meeting notes).",
          must_cite=[S(file=prd, page=1), S(file="meeting_notes_2025-11-03.txt")],
          expected_keywords=["Q3 2026", "battery"]),
        G(id="xd-03", type="cross_doc",
          question="Who supplies Orion's battery cells and what is that contract worth?",
          reference="Kestrel Components GmbH; USD 1,250,000 over 24 months.",
          must_cite=[S(file="supplier_agreement_scanned.pdf", page=1)],
          expected_keywords=["Kestrel"], expected_numbers=["1250000"]),
        # summary
        G(id="sum-01", type="summary", question="Summarise the key decisions from the 2025-11-03 meeting.",
          reference="Orion v2 moved to Q3 2026; CI to self-hosted runners by end of Jan 2026; "
                    "Vega defect detection as paid add-on.",
          must_cite=[S(file="meeting_notes_2025-11-03.txt")], expected_keywords=["Q3 2026"]),
        # follow-ups
        G(id="fu-01", type="followup", history=["What was revenue by region in FY2025?"],
          question="And which of those regions grew the slowest?",
          reference="LATAM, +3.2%.", must_cite=[S(file=rep, page=2)], expected_keywords=["LATAM"]),
        G(id="fu-02", type="followup", history=["How many support tickets are in the log?"],
          question="How many of them are critical?", reference=str(n_crit),
          must_cite=[S(file=csv)], expected_numbers=[str(n_crit)]),
        # unanswerable
        G(id="na-01", type="unanswerable", answerable=False,
          question="What is Acme Robotics' stock ticker symbol?", reference="Not in the documents."),
        G(id="na-02", type="unanswerable", answerable=False,
          question="Who is Acme's chief financial officer?", reference="Not in the documents."),
        G(id="na-03", type="unanswerable", answerable=False,
          question="What was Acme's revenue in FY2022?", reference="Not in the documents."),
    ]  # fmt: skip
    if needle_page:
        items.append(G(
            id="lg-01", type="lookup", requires_large=True,
            question="What is the daily meal allowance for business trips to Tokyo?",
            reference="JPY 9,500.", must_cite=[S(file="acme_policies_large.pdf", page=needle_page)],
            expected_numbers=["9500"],
        ))  # fmt: skip
    return items


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawTextHelpFormatter)
    ap.add_argument("--out", type=Path, default=Path("demo_corpus"))
    ap.add_argument("--large-pages", type=int, default=0, help="also write an N-page policy PDF")
    args = ap.parse_args()
    out: Path = args.out
    out.mkdir(parents=True, exist_ok=True)

    rng = random.Random(SEED)
    orders, tickets = build_orders(rng), build_tickets(rng)
    actual = orders.groupby("region")["revenue"].sum()
    targets = {r: int(round(actual[r] * TARGET_FACTORS[r], -4)) for r in REGIONS}

    with tempfile.TemporaryDirectory() as tmp:
        make_annual_report(out / "acme_annual_report_2025.pdf", Path(tmp))
    make_scanned_agreement(out / "supplier_agreement_scanned.pdf")
    make_invoice_scan(out / "invoice_scan.png")
    make_active_robots_chart(out / "monthly_active_robots.png")
    make_dashboard(out / "ops_dashboard.png")
    make_prd(out / "orion_amr_v2_prd.docx")
    make_board_deck(out / "board_update_q4.pptx")
    make_handbook(out / "engineering_handbook.md")
    make_meeting_notes(out / "meeting_notes_2025-11-03.txt")
    make_code(out / "inventory_service.py")
    write_sales_workbook(out / "sales_2025.xlsx", orders, targets)
    tickets.to_csv(out / "support_tickets.csv", index=False)
    needle = make_large_policy_pdf(out / "acme_policies_large.pdf", args.large_pages) \
        if args.large_pages else None  # fmt: skip

    golden = build_golden(orders, tickets, targets, needle)
    write_golden(out / "golden.jsonl", golden)
    manifest = {"seed": SEED, "files": sorted(p.name for p in out.iterdir() if p.is_file())}
    (out / "manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    print(f"Wrote {len(manifest['files'])} files and {len(golden)} golden questions to {out}")


if __name__ == "__main__":
    main()
